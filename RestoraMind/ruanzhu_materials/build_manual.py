# -*- coding: utf-8 -*-
"""把说明书内容模型渲染为 PDF（reportlab）与 Word（python-docx）。"""
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus.tableofcontents import TableOfContents
from reportlab.platypus import (CondPageBreak, BaseDocTemplate, Frame, Image, PageBreak, PageTemplate, Paragraph, Preformatted,
                                Spacer, Table, TableStyle, KeepTogether)

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import manual_content  # noqa: E402
from restoramind import __copyright_holder__, __software_name__, __version__  # noqa: E402

FONT = "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc"
TITLE = f"{__software_name__} V{__version__}"
DOC_NAME = "软件说明书"
OUT_PDF = HERE / "manual_V1.0.pdf"
OUT_DOCX = HERE / "manual_V1.0.docx"


import unicodedata


def _w(ch):
    return 2 if unicodedata.east_asian_width(ch) in "WF" else 1


def wrap_code(text, limit):
    """按显示宽度折行（优先在空格处断开），续行缩进两格，避免代码块溢出页面。"""
    out = []
    for line in text.split("\n"):
        cur = line
        first = True
        while sum(_w(c) for c in cur) > limit:
            width, cut, last_space = 0, 0, -1
            for i, ch in enumerate(cur):
                if width + _w(ch) > limit:
                    break
                width += _w(ch)
                cut = i + 1
                if ch == " " and i > 0:
                    last_space = i
            if last_space > limit * 0.4:
                cut = last_space
            out.append(cur[:cut].rstrip())
            cur = "  " + cur[cut:].lstrip()
            first = False
        out.append(cur)
    return "\n".join(out)


# ============================================================ PDF
def esc(t: str) -> str:
    return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def build_pdf(blocks):
    pdfmetrics.registerFont(TTFont("CJK", FONT, subfontIndex=0))
    st = {
        "p": ParagraphStyle("p", fontName="CJK", fontSize=10.5, leading=17, firstLineIndent=21, wordWrap="CJK", spaceAfter=3),
        "h1": ParagraphStyle("h1", fontName="CJK", fontSize=17, leading=24, spaceBefore=6, spaceAfter=10, wordWrap="CJK"),
        "h2": ParagraphStyle("h2", fontName="CJK", fontSize=13.5, leading=20, spaceBefore=8, spaceAfter=5, wordWrap="CJK"),
        "h3": ParagraphStyle("h3", fontName="CJK", fontSize=11, leading=16, spaceBefore=6, spaceAfter=3, wordWrap="CJK"),
        "cell": ParagraphStyle("cell", fontName="CJK", fontSize=8.8, leading=12.4, wordWrap="CJK"),
        "head": ParagraphStyle("head", fontName="CJK", fontSize=9, leading=12.5, wordWrap="CJK", textColor=colors.black),
        "bullet": ParagraphStyle("b", fontName="CJK", fontSize=10.5, leading=17, leftIndent=18, bulletIndent=6, wordWrap="CJK", spaceAfter=2),
        "cap": ParagraphStyle("cap", fontName="CJK", fontSize=9.5, leading=14, alignment=TA_CENTER, spaceAfter=6, wordWrap="CJK"),
        "code": ParagraphStyle("code", fontName="CJK", fontSize=8.4, leading=11.6, leftIndent=6, backColor=colors.HexColor("#f4f4f4"),
                               borderPadding=4, spaceBefore=4, spaceAfter=6),
        "cover_t": ParagraphStyle("ct", fontName="CJK", fontSize=24, leading=34, alignment=TA_CENTER, wordWrap="CJK"),
        "cover": ParagraphStyle("cv", fontName="CJK", fontSize=14, leading=26, alignment=TA_CENTER, wordWrap="CJK"),
    }
    width, height = A4
    left = right = 2.2 * cm
    frame_w = width - left - right

    def decorate(c, doc):
        c.saveState()
        c.setFont("CJK", 9.5)
        c.drawString(left, height - 1.4 * cm, f"{TITLE} {DOC_NAME}")
        c.drawRightString(width - right, height - 1.4 * cm, f"第 {doc.page} 页")
        c.setLineWidth(0.5)
        c.line(left, height - 1.6 * cm, width - right, height - 1.6 * cm)
        c.restoreState()

    class ManualDoc(BaseDocTemplate):
        def afterFlowable(self, flowable):
            style = getattr(flowable, "style", None)
            if isinstance(flowable, Paragraph) and style is not None and style.name in ("h1", "h2"):
                self.notify("TOCEntry", (0 if style.name == "h1" else 1, flowable.getPlainText(), self.page))

    doc = ManualDoc(str(OUT_PDF), pagesize=A4, leftMargin=left, rightMargin=right, topMargin=2.2 * cm, bottomMargin=2.2 * cm,
                          title=f"{TITLE} {DOC_NAME}")
    doc.addPageTemplates([PageTemplate("main", [Frame(left, 2.2 * cm, frame_w, height - 4.4 * cm, id="f")], onPage=decorate)])

    st["cover_t"].fontSize, st["cover_t"].leading = 18, 26
    st["cover"].fontSize, st["cover"].leading = 11.5, 18
    story = [Paragraph(esc(__software_name__), st["cover_t"]),
             Paragraph(f"版本：V{__version__}　{DOC_NAME}　著作权人：{esc(__copyright_holder__)}", st["cover"]), Spacer(1, 0.4 * cm)]

    # 目录由标题自动生成（简版：章标题列表）
    toc = TableOfContents(dotsMinLevel=0)
    toc.levelStyles = [
        ParagraphStyle("toc1", fontName="CJK", fontSize=10.5, leading=15, leftIndent=0, spaceBefore=3, wordWrap="CJK"),
        ParagraphStyle("toc2", fontName="CJK", fontSize=9.5, leading=13, leftIndent=18, wordWrap="CJK"),
    ]
    story.append(Paragraph("目录", ParagraphStyle("toc_title", parent=st["h2"])))
    story.append(toc)
    story.append(Spacer(1, 0.3 * cm))

    first_h1 = True
    for blk in blocks:
        kind = blk[0]
        if kind == "h1":
            first_h1 = False
            story.append(CondPageBreak(4.5 * cm))
            story.append(Paragraph(esc(blk[1]), st["h1"]))
        elif kind in ("h2", "h3"):
            story.append(CondPageBreak(3.2 * cm))
            story.append(Paragraph(esc(blk[1]), st[kind]))
        elif kind == "p":
            story.append(Paragraph(esc(blk[1]), st["p"]))
        elif kind == "bullets":
            for item in blk[1]:
                story.append(Paragraph(esc(item), st["bullet"], bulletText="•"))
        elif kind == "numbered":
            for i, item in enumerate(blk[1], 1):
                story.append(Paragraph(esc(item), st["bullet"], bulletText=f"{i}."))
        elif kind == "code":
            story.append(Preformatted(wrap_code(blk[1], 100), st["code"]))
        elif kind == "fig":
            from PIL import Image as PILImage
            w, h = PILImage.open(blk[1]).size
            iw = min(frame_w, 14.0 * cm, 6.2 * cm * w / h)
            story.append(KeepTogether([Image(str(blk[1]), width=iw, height=iw * h / w), Paragraph(esc(blk[2]), st["cap"])]))
        elif kind == "table":
            header, rows, ratios = blk[1], blk[2], blk[3]
            data = [[Paragraph(esc(h), st["head"]) for h in header]]
            data += [[Paragraph(esc(str(c)), st["cell"]) for c in r] for r in rows]
            t = Table(data, colWidths=[frame_w * r for r in ratios], repeatRows=1)
            t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#888888")),
                                   ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e6e6e6")),
                                   ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                   ("TOPPADDING", (0, 0), (-1, -1), 2.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5)]))
            t.spaceAfter = 6
            story.append(t)
        elif kind == "pagebreak":
            story.append(PageBreak())
    doc.multiBuild(story)


# ============================================================ DOCX
def set_run_font(run, size=None, bold=None, mono=False):
    name = "Courier New" if mono else "Times New Roman"
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def add_field(paragraph, instr):
    run = paragraph.add_run()
    for tag, text in (("begin", None), (None, instr), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    set_run_font(run, 10)


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def build_docx(blocks):
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = sec.bottom_margin = Cm(2.2)
    normal = d.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.25

    # 页眉：软件名称+版本 / 右上角页码；不设页脚
    hp = sec.header.paragraphs[0]
    hp.paragraph_format.tab_stops.add_tab_stop(Cm(16.6), WD_TAB_ALIGNMENT.RIGHT)
    r = hp.add_run(f"{TITLE} {DOC_NAME}\t第 "); set_run_font(r, 10)
    add_field(hp, "PAGE")
    r = hp.add_run(" 页"); set_run_font(r, 10)

    def para(text, size=10.5, bold=False, align=None, indent=True, space_after=3, mono=False):
        p = d.add_paragraph()
        if align:
            p.alignment = align
        if indent:
            p.paragraph_format.first_line_indent = Pt(21)
        p.paragraph_format.space_after = Pt(space_after)
        run = p.add_run(text)
        set_run_font(run, size, bold, mono)
        return p

    para(__software_name__, 18, True, WD_ALIGN_PARAGRAPH.CENTER, indent=False, space_after=6)
    para(f"版本：V{__version__}　{DOC_NAME}　著作权人：{__copyright_holder__}", 11.5, False, WD_ALIGN_PARAGRAPH.CENTER,
         indent=False, space_after=8)
    para("目录", 13.5, True, indent=False, space_after=4)
    # 目录域：打开文档时由 Word 生成带页码的目录（settings 中已开启打开时更新域）
    tp = d.add_paragraph()
    run = tp.add_run()
    for tag, text in (("begin", None), (None, 'TOC \\o "1-2" \\h \\z \\u'), ("separate", None)):
        if tag:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text
        run._r.append(el)
    set_run_font(tp.add_run("（目录将在打开文档时自动生成；如未显示，请在目录处右键选择“更新域”）"), 9)
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    tp.add_run()._r.append(end)
    upd = OxmlElement("w:updateFields"); upd.set(qn("w:val"), "true")
    d.settings.element.append(upd)
    cp = d.core_properties
    cp.author, cp.title, cp.subject = __copyright_holder__, f"{TITLE} {DOC_NAME}", __software_name__
    cp.last_modified_by, cp.comments = __copyright_holder__, ""

    first_h1 = True
    sizes = {"h1": 17, "h2": 13.5, "h3": 11}
    for blk in blocks:
        kind = blk[0]
        if kind == "h1":
            first_h1 = False
        if kind in sizes:
            p = para(blk[1], sizes[kind], True, indent=False, space_after=6)
            p.paragraph_format.space_before = Pt(8 if kind != "h1" else 4)
            p.paragraph_format.keep_with_next = True
            p.style = d.styles["Heading %s" % kind[1]]
            for run in p.runs:
                set_run_font(run, sizes[kind], True)
                run.font.color.rgb = RGBColor(0, 0, 0)
        elif kind == "p":
            para(blk[1])
        elif kind == "bullets":
            for item in blk[1]:
                p = para("• " + item, indent=False)
                p.paragraph_format.left_indent = Pt(18)
        elif kind == "numbered":
            for i, item in enumerate(blk[1], 1):
                p = para(f"{i}. {item}", indent=False)
                p.paragraph_format.left_indent = Pt(18)
        elif kind == "code":
            for line in wrap_code(blk[1], 88).split("\n"):
                p = para(line if line else " ", 8.5, indent=False, space_after=0, mono=True)
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.left_indent = Pt(10)
        elif kind == "fig":
            from PIL import Image as _PI
            _w, _h = _PI.open(blk[1]).size
            d.add_picture(str(blk[1]), width=Cm(min(14.0, 6.2 * _w / _h)))
            d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            para(blk[2], 9.5, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, space_after=8)
        elif kind == "table":
            header, rows, ratios = blk[1], blk[2], blk[3]
            t = d.add_table(rows=1, cols=len(header))
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            total = 16.6
            for i, h in enumerate(header):
                c = t.rows[0].cells[i]
                c.width = Cm(total * ratios[i])
                c.paragraphs[0].text = ""
                set_run_font(c.paragraphs[0].add_run(h), 9, True)
                shade(c, "E6E6E6")
            for r_ in rows:
                cells = t.add_row().cells
                for i, v in enumerate(r_):
                    cells[i].width = Cm(total * ratios[i])
                    set_run_font(cells[i].paragraphs[0].add_run(str(v)), 9)
            d.add_paragraph().paragraph_format.space_after = Pt(2)
    d.save(str(OUT_DOCX))


if __name__ == "__main__":
    blocks = manual_content.build()
    build_pdf(blocks)
    build_docx(blocks)
    import pymupdf
    n = len(pymupdf.open(str(OUT_PDF)))
    print(f"PDF 页数={n}  -> {OUT_PDF.name}\nWord -> {OUT_DOCX.name}")
